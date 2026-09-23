# services/export_service.py
"""Excel 导出（xlsx）"""
import time
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


_STATUS_LABEL = {
    'pending':   '待付款',
    'paid':      '已付款',
    'cooking':   '制作中',
    'done':      '已完成',
    'cancelled': '已取消',
}


def _style_header(ws, ncol):
    fill = PatternFill('solid', fgColor='EEF2F7')
    font = Font(bold=True, color='1F2937', size=11)
    align = Alignment(horizontal='center', vertical='center')
    for i in range(1, ncol + 1):
        c = ws.cell(row=1, column=i)
        c.fill = fill
        c.font = font
        c.alignment = align
    ws.freeze_panes = 'A2'
    ws.row_dimensions[1].height = 22


def _auto_width(ws, rows, min_w=8, max_w=40):
    if not rows:
        return
    ncol = len(rows[0])
    for i in range(1, ncol + 1):
        maxlen = 0
        for r in rows:
            v = r[i - 1] if i - 1 < len(r) else None
            if v is None:
                continue
            s = str(v)
            w = sum(2 if ord(ch) > 127 else 1 for ch in s)
            if w > maxlen:
                maxlen = w
        width = min(max(maxlen + 2, min_w), max_w)
        ws.column_dimensions[get_column_letter(i)].width = width


def _save(wb):
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _today():
    return time.strftime('%Y-%m-%d')


def orders_xlsx():
    from repositories.instances import order_repo
    orders = order_repo.get_all()
    orders.sort(key=lambda o: o.get('created_at') or '', reverse=True)

    wb = Workbook()
    ws = wb.active
    ws.title = '订单明细'

    headers = ['订单号', '用户UID', '桌号', '菜品', '数量合计', '金额(元)', '状态', '下单时间']
    ws.append(headers)

    summary = wb.create_sheet('订单汇总')
    summary.append(['日期', '订单数', '营业额(元)'])

    by_date = {}
    for o in orders:
        ct = (o.get('created_at') or '')[:10]
        if o.get('status') == 'cancelled':
            continue
        d = by_date.setdefault(ct, {'cnt': 0, 'total': 0.0})
        d['cnt'] += 1
        d['total'] += float(o.get('total', 0) or 0)

    total_amount = 0.0
    total_qty = 0
    for o in orders:
        items = o.get('items') or []
        items_str = '、'.join(f"{it.get('name')}x{it.get('qty')}" for it in items)
        qty_sum = sum(int(it.get('qty', 1)) for it in items)
        amt = float(o.get('total', 0) or 0)
        if o.get('status') != 'cancelled':
            total_amount += amt
            total_qty += qty_sum
        ws.append([
            o.get('id', ''),
            o.get('uid') or '匿名',
            o.get('table_no') or '外带',
            items_str,
            qty_sum,
            round(amt, 2),
            _STATUS_LABEL.get(o.get('status'), o.get('status', '')),
            o.get('created_at', ''),
        ])

    for d in sorted(by_date.keys()):
        v = by_date[d]
        summary.append([d, v['cnt'], round(v['total'], 2)])

    _style_header(ws, len(headers))
    _style_header(summary, 3)

    for row in ws.iter_rows(min_row=2, min_col=6, max_col=6):
        for c in row:
            c.number_format = '#,##0.00'
    for row in summary.iter_rows(min_row=2, min_col=3, max_col=3):
        for c in row:
            c.number_format = '#,##0.00'

    _auto_width(ws, list(ws.iter_rows(values_only=True)))
    _auto_width(summary, list(summary.iter_rows(values_only=True)))

    ws.append([])
    ws.append(['合计（不含已取消）', '', '', '', total_qty, round(total_amount, 2), '', ''])

    return _save(wb), f'orders_{_today()}.xlsx'


def users_xlsx():
    from repositories.instances import member_repo
    users = member_repo.get_all()
    users.sort(key=lambda u: u.get('created_at') or '', reverse=True)

    wb = Workbook()
    ws = wb.active
    ws.title = '用户'

    headers = ['UID', '用户名', '昵称', '积分', '等级', '状态', '注册时间', '最近积分变动']
    ws.append(headers)

    for u in users:
        ws.append([
            u.get('uid', ''),
            u.get('username', ''),
            u.get('nickname', ''),
            int(u.get('points', 0) or 0),
            u.get('level', ''),
            '已禁用' if u.get('disabled') else '正常',
            u.get('created_at', ''),
            u.get('last_points_at', ''),
        ])

    _style_header(ws, len(headers))
    _auto_width(ws, list(ws.iter_rows(values_only=True)))
    return _save(wb), f'users_{_today()}.xlsx'


def points_xlsx():
    from repositories.instances import points_log_repo, member_repo
    logs = points_log_repo.get_all()
    logs.sort(key=lambda x: x.get('created_at') or '', reverse=True)

    users = {u['uid']: u.get('username') or u.get('nickname') or ''
             for u in member_repo.get_all()}

    wb = Workbook()
    ws = wb.active
    ws.title = '积分流水'

    headers = ['时间', 'UID', '用户', '变动', '原因']
    ws.append(headers)

    for lg in logs:
        ws.append([
            lg.get('created_at', ''),
            lg.get('uid', ''),
            users.get(lg.get('uid'), ''),
            int(lg.get('change', 0) or 0),
            lg.get('reason', ''),
        ])

    _style_header(ws, len(headers))
    _auto_width(ws, list(ws.iter_rows(values_only=True)))
    return _save(wb), f'points_{_today()}.xlsx'


class _ExportService:
    orders = staticmethod(orders_xlsx)
    users = staticmethod(users_xlsx)
    points = staticmethod(points_xlsx)


export_service = _ExportService()
